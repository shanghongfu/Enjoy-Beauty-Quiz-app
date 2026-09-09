# -*- coding: utf-8 -*-
"""生成《Enjoy Beauty 刷题 App 商用竞品分析与差距报告》docx"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

ACCENT = RGBColor(0x18, 0x5F, 0xA5)   # 蓝
GREEN  = RGBColor(0x0F, 0x6E, 0x56)
RED    = RGBColor(0xA3, 0x2D, 0x2D)
GRAY   = RGBColor(0x5F, 0x5E, 0x5A)
DARK   = RGBColor(0x2C, 0x2C, 0x2A)

def set_font(run, size=10.5, bold=False, color=None, name='微软雅黑'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    r = run._element.rPr.rFonts
    r.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = color

def add_h1(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=16, bold=True, color=ACCENT)
    p.space_after = Pt(6)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=13, bold=True, color=DARK)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_body(doc, text, size=10.5, color=None, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    return p

def add_bullet(doc, text, color=None, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r0 = p.add_run(bold_prefix)
        set_font(r0, size=10.5, bold=True, color=color)
    r = p.add_run(text)
    set_font(r, size=10.5, color=color)
    p.paragraph_format.space_after = Pt(2)
    return p

def style_table(tbl):
    tbl.style = 'Light Grid Accent 1'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

def fill_cell(cell, text, bold=False, color=None, size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)

doc = Document()
# 页面边距
for s in doc.sections:
    s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

# 封面标题
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Enjoy Beauty 刷题 App\n商用化竞品分析与差距报告')
set_font(r, size=22, bold=True, color=ACCENT)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('对标国内外同类产品 · 找出商用化短板 · 给出落地路线')
set_font(r, size=12, color=GRAY)
doc.add_paragraph()

# 一、项目现状
add_h1(doc, '一、产品现状')
add_body(doc, 'Enjoy Beauty 是一款面向美容/美业职业技能考试的刷题练习应用，当前以 PWA 形态部署在公网链接，支持中英双语。核心能力：')
for t in ['多题库管理 + Word 批量导入（近期已修复 docx 解析 bug，可全量识别 465 题大题库）',
          '单选 / 多选 / 判断三类题型',
          '顺序 / 随机 / 错题重练 / 模拟考试四种练习模式',
          '错题本 + 答题统计 + 连续打卡（streak）',
          '管理员 / 学员账号体系 + 题库分配（学员只见被分配的题库）',
          'PWA 离线可用，移动端优先']:
    add_bullet(doc, t)

# 二、竞品情报
add_h1(doc, '二、同类产品调研')
add_h2(doc, '2.1 国内竞品（美容师职业资格考证赛道）')
tbl = doc.add_table(rows=1, cols=5)
style_table(tbl)
hdr = ['产品', '题库规模', '核心功能', '付费模式', '突出特点']
for i, h in enumerate(hdr):
    fill_cell(tbl.rows[0].cells[i], h, bold=True, color=ACCENT)
rows = [
    ['美容师考试聚题库', '初/中/高三级全考点', '专项练习、随机刷题、全真模考、错题复盘、搜题(文字/拍照/语音)', '基础免费+付费会员', '按知识点+分值占比拆题，可视化学习报告'],
    ['美容师题库(希赛)', '初级/中级/高级', '每日一练、章节练习、学习计划、智能出题、拍照搜题、考试指南', '免费+会员解锁', '学习计划自定义+进度跟踪'],
    ['美容师题库(多款)', '覆盖考纲全部考点', '刷题/拍照搜题/错题复盘/模拟考场', '免费基础+付费', '离线缓存、签到免费领模考次数'],
    ['优题宝/上学吧', '全等级适配', '分章节专项、模拟卷、三种搜题', '订阅制', '题库按考纲板块划分'],
]
for row in rows:
    cells = tbl.add_row().cells
    for i, v in enumerate(row):
        fill_cell(cells[i], v)

add_h2(doc, '2.2 海外竞品（NIC 美容师执照考试赛道）')
tbl2 = doc.add_table(rows=1, cols=5)
style_table(tbl2)
for i, h in enumerate(hdr):
    fill_cell(tbl2.rows[0].cells[i], h, bold=True, color=ACCENT)
rows2 = [
    ['NIC Exam Prep', '按官方 CIB 比重', '题库+按域闪卡+限时模考(90分钟)+全解析(含错误项为什么错)', '一次买断终身', '免费10题试用无需注册，Audited 题库'],
    ['GlowStudy (Esthetician State Board)', '2450题(1450多选+1000判断)', '7种学习模式(Daily Glow/智能复习/速答/闪卡/错题)+4级难度', '订阅制(周$4.99/月$8.99/年$34.99)', '认知科学(置信度校准+间隔重复)，Exam Ready 证书+保过退款'],
    ['Prepdex', '1188+题/10主题', '间隔重复+限时模考+AI解析+备考评分+Pass Plan', 'Pass Plan 订阅', '诊断报告定位薄弱域，通过保证'],
    ['Esthetician State Board Prep', 'NIC 内容自研', '限时模考+每题解析+进度本地存储', '一次性买断', '无账号、无广告、离线'],
]
for row in rows2:
    cells = tbl2.add_row().cells
    for i, v in enumerate(row):
        fill_cell(cells[i], v)

add_h2(doc, '2.3 行业盈利模式')
add_bullet(doc, '免费基础 + 付费会员订阅（月/季/年，年费约 30-300 元，海外 35-100 美元），高频刚需功能(解析、模考、错题)设为付费墙', bold_prefix='会员订阅制：')
add_bullet(doc, '基础免费 + 开屏/信息流/结果页广告，或激励视频(看广告解锁次数)', bold_prefix='广告变现：')
add_bullet(doc, '单套真题/押题/冲刺包 9.9-49.9 元买断', bold_prefix='题库售卖：')
add_bullet(doc, '面向培训机构/美容学校的定制考试平台、按学生数或年费收取，教师端看学情', bold_prefix='B端SaaS：')
add_bullet(doc, '与教培机构 CPS 分成(报名课程佣金20%-50%)', bold_prefix='导流分成：')

# 三、差距分析
add_h1(doc, '三、差距分析：咱们 vs 竞品')
add_body(doc, '对比竞品，Enjoy Beauty 已经具备“能跑通的刷题闭环”，但在决定商业转化和留存的功能上有明显短板。', color=DARK)

tbl3 = doc.add_table(rows=1, cols=4)
style_table(tbl3)
hdr3 = ['能力维度', '咱们现状', '竞品标配', '商用差距']
for i, h in enumerate(hdr3):
    fill_cell(tbl3.rows[0].cells[i], h, bold=True, color=ACCENT)
rows3 = [
    ['题库管理', '✅ 多库+Word导入', '✅ 支持docx/xlsx/pdf批量导入', '✅ 已达，可加强图片/公式支持'],
    ['错题本', '✅ 有基础版', '✅ 自动归类+按知识点分类+导出组卷+连续答对移出', '⚠️ 缺分类、缺导出、缺智能重练'],
    ['答题解析', '❌ 导入解析字段有，但界面展示弱', '✅ 每题完整解析(含错误项解释)', '❌ 内容与展示都缺，核心短板'],
    ['模拟考场', '⚠️ 有exam模式(50题)但无计时', '✅ 限时+题型比重+仿真机考', '❌ 无计时、无组卷比例、无报告'],
    ['学习报告', '⚠️ 有统计页', '✅ 可视化报告+薄弱点定位+趋势', '⚠️ 缺薄弱点定位和趋势'],
    ['搜题', '❌ 无', '✅ 文字/拍照/语音三合一', '❌ 完全缺失'],
    ['学习计划/打卡', '⚠️ 有streak', '✅ 每日一练/任务/计划/签到', '⚠️ 打卡简单，缺任务体系'],
    ['排行榜/班级', '❌ 无', '✅ 班级+排行榜+教师端', '❌ 完全缺失(B端入口)'],
    ['商业化', '❌ 无支付无会员', '✅ 会员/买断/广告/题库售卖', '❌ 完全缺失'],
]
for row in rows3:
    cells = tbl3.add_row().cells
    for i, v in enumerate(row):
        fill_cell(cells[i], v, color=(RED if v.startswith('❌') else GREEN if v.startswith('✅') else GRAY if v.startswith('⚠️') else None))

# 四、分优先级落地路线
add_h1(doc, '四、商用化落地路线（按优先级）')

add_h2(doc, 'P0 · 根基加固（1-2 周）')
for t in ['题目解析能力：导入时识别“解析/Explanation”行入库；答题后展示解析（含错误项为什么错）',
          '数据安全：补上数据备份/导出能力，题库是核心资产，防止误删',
          '合规声明：页脚加“非官方、仅供学习参考”声明，避免与官方/版权冲突']:
    add_bullet(doc, t)

add_h2(doc, 'P1 · 学习闭环（2-4 周，决定留存）')
for t in ['模拟考场升级：限时(按真实考试时长)、按知识点/章节比重组卷、交卷生成成绩报告',
          '错题本升级：按知识点分类、错题重练、连续答对2次自动移出、导出错题组卷',
          '学习报告：答题趋势折线、各知识点掌握度、薄弱点定位(参考 Prepdex/GlowStudy)',
          '每日一练 + 签到打卡：提升日活与粘性']:
    add_bullet(doc, t)

add_h2(doc, 'P2 · 商业闭环（4-8 周，开始赚钱）')
for t in ['支付 + 会员体系：微信/支付宝支付，免费基础 + 付费解锁(解析/模考/完整题库)',
          '题库售卖/订阅：单套题库买断或年卡订阅，参考 9.9-49.9 元单套、30-300 元/年',
          '搜题能力：先做关键词搜索(简单)，拍照/语音搜题(需接 OCR/ASR，后期)']:
    add_bullet(doc, t)

add_h2(doc, 'P3 · 增长与 B 端（8 周后）')
for t in ['排行榜 + 班级 + 教师端：教师可建班、布置练习、看学情报告 → B 端机构采购入口',
          '导流分成：与美容培训机构 CPS 合作',
          '题库 UGC：用户可上传题库，付费下载分成(参考刷刷题 10-20% 佣金)']:
    add_bullet(doc, t)

# 五、关键建议
add_h1(doc, '五、最关键的三个判断')
for t in ['差异化定位想清楚：咱们的独特价值是“机构/老师可自建题库、学员扫码即用、中英双语”，这是通用题库 App 做不到的，建议主打 B 端机构+SaaS 模式，而非和 C 端题库 App 拼题量。',
          '题库内容合规是红线：美容师题库涉及版权(真题/解析)，务必用自研题或授权题，避免商用后被追责。这也是为什么“机构自带题库”模式更有壁垒——题是客户自己的，没有版权风险。',
          '先跑通付费闭环再谈规模：哪怕只有 1 个付费功能(如解锁解析/模考)也要先上线收钱验证需求，别一次性铺所有功能。']:
    add_bullet(doc, t, color=DARK)

# 尾注
doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('— 报告完 —')
set_font(r, size=10, color=GRAY)

out = '/Users/ak/WorkBuddy/2026-08-25-13-22-56/quiz-app/output/EnjoyBeauty刷题App商用竞品分析与差距报告.docx'
doc.save(out)
print('已生成:', out)
